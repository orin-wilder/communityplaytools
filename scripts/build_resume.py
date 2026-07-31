from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "assets" / "resume"
OUT_DIR.mkdir(parents=True, exist_ok=True)
OUT = OUT_DIR / "Ryan_Stock_Resume.docx"

PINE = "103F38"
CORAL = "B63F28"
INK = "1A1714"
MUTED = "4A443D"
LINE = "D8D1C8"
PAPER = "F5F1ED"

# Compact_reference_guide base, with a named resume_compact_two_page override:
# Letter portrait; margins 0.62" L/R and 0.55" T/B; Arial 9.7 pt body;
# Georgia display headings; real bullet numbering at 0.16"/0.31"; no tables/columns.

doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.55)
section.bottom_margin = Inches(0.55)
section.left_margin = Inches(0.62)
section.right_margin = Inches(0.62)
section.header_distance = Inches(0.25)
section.footer_distance = Inches(0.27)

styles = doc.styles


def set_font(run, name="Arial", size=9.7, color=INK, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


normal = styles["Normal"]
normal.font.name = "Arial"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
normal.font.size = Pt(9.7)
normal.font.color.rgb = RGBColor.from_string(INK)
normal.paragraph_format.space_before = Pt(0)
normal.paragraph_format.space_after = Pt(2.6)
normal.paragraph_format.line_spacing = 1.04

for style_name, size in [("Heading 1", 11.4), ("Heading 2", 10.25), ("Heading 3", 9.6)]:
    style = styles[style_name]
    style.font.name = "Georgia"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Georgia")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Georgia")
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor.from_string(PINE)
    style.paragraph_format.keep_with_next = True
    style.paragraph_format.space_before = Pt(6 if style_name == "Heading 1" else 3.5)
    style.paragraph_format.space_after = Pt(2.5)
    style.paragraph_format.line_spacing = 1.04


def add_real_bullet_definition(document):
    numbering = document.part.numbering_part.element
    abstract_ids = [
        int(node.get(qn("w:abstractNumId")))
        for node in numbering.findall(qn("w:abstractNum"))
    ]
    num_ids = [int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=0) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "bullet")
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), "•")
    lvl_jc = OxmlElement("w:lvlJc")
    lvl_jc.set(qn("w:val"), "left")
    ppr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "446")
    tabs.append(tab)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "446")
    ind.set(qn("w:hanging"), "216")
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:before"), "0")
    spacing.set(qn("w:after"), "44")
    spacing.set(qn("w:line"), "250")
    spacing.set(qn("w:lineRule"), "auto")
    ppr.extend([tabs, ind, spacing])
    rpr = OxmlElement("w:rPr")
    rfonts = OxmlElement("w:rFonts")
    rfonts.set(qn("w:ascii"), "Arial")
    rfonts.set(qn("w:hAnsi"), "Arial")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), CORAL)
    rpr.extend([rfonts, color])
    lvl.extend([start, num_fmt, lvl_text, lvl_jc, ppr, rpr])
    abstract.append(lvl)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


BULLET_NUM_ID = add_real_bullet_definition(doc)


def add_hyperlink(paragraph, text, url, color=PINE, underline=True):
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    rfonts = OxmlElement("w:rFonts")
    rfonts.set(qn("w:ascii"), "Arial")
    rfonts.set(qn("w:hAnsi"), "Arial")
    rpr.append(rfonts)
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "17")
    rpr.append(size)
    c = OxmlElement("w:color")
    c.set(qn("w:val"), color)
    rpr.append(c)
    if underline:
        u = OxmlElement("w:u")
        u.set(qn("w:val"), "single")
        rpr.append(u)
    run.append(rpr)
    text_el = OxmlElement("w:t")
    text_el.text = text
    run.append(text_el)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_section_heading(text):
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(6.5)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.04
    r = p.add_run(text.upper())
    set_font(r, "Arial", 9.35, PINE, bold=True)
    r.font.letter_spacing = Pt(0.6)
    return p


def add_role(title, company, dates, location=None, descriptor=None):
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(3.7)
    p.paragraph_format.space_after = Pt(0.8)
    p.paragraph_format.line_spacing = 1.04
    r = p.add_run(title)
    set_font(r, "Arial", 9.95, INK, bold=True)
    r = p.add_run(f"  |  {company}")
    set_font(r, "Arial", 9.95, PINE, bold=True)
    r = p.add_run(f"  |  {dates}")
    set_font(r, "Arial", 9.55, MUTED, bold=False)
    if location or descriptor:
        p2 = doc.add_paragraph()
        p2.paragraph_format.keep_with_next = True
        p2.paragraph_format.space_before = Pt(0)
        p2.paragraph_format.space_after = Pt(1.5)
        p2.paragraph_format.line_spacing = 1.04
        r2 = p2.add_run("  |  ".join(x for x in [location, descriptor] if x))
        set_font(r2, "Arial", 8.9, MUTED, italic=True)


def add_subrole(text):
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(2.8)
    p.paragraph_format.space_after = Pt(0.8)
    r = p.add_run(text.upper())
    set_font(r, "Arial", 8.55, CORAL, bold=True)


def add_bullet(text, keep_next=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2.2)
    p.paragraph_format.line_spacing = 1.04
    p.paragraph_format.keep_together = True
    p.paragraph_format.keep_with_next = keep_next
    ppr = p._p.get_or_add_pPr()
    numpr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    numid = OxmlElement("w:numId")
    numid.set(qn("w:val"), str(BULLET_NUM_ID))
    numpr.extend([ilvl, numid])
    ppr.append(numpr)
    r = p.add_run(text)
    set_font(r, "Arial", 9.6, INK)
    return p


def add_body(text, size=9.7, color=INK, bold_prefix=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2.8)
    p.paragraph_format.line_spacing = 1.04
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        set_font(r, "Arial", size, color, bold=True)
        r = p.add_run(text[len(bold_prefix):])
        set_font(r, "Arial", size, color)
    else:
        r = p.add_run(text)
        set_font(r, "Arial", size, color)
    return p


# First-page identity block, based on a restrained customer_pack pattern.
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(0)
p.paragraph_format.space_after = Pt(0)
r = p.add_run("RYAN STOCK")
set_font(r, "Georgia", 27, INK, bold=True)

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(0)
p.paragraph_format.space_after = Pt(2.5)
r = p.add_run("Senior Program & Innovation Leader")
set_font(r, "Georgia", 13.4, PINE, bold=True)

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(0)
p.paragraph_format.space_after = Pt(4.2)
p.paragraph_format.line_spacing = 1.0
for i, (label, url) in enumerate([
    ("St. Petersburg, FL", None),
    ("858.414.8819", "tel:+18584148819"),
    ("ryan@communityplaytools.com", "mailto:ryan@communityplaytools.com"),
    ("linkedin.com/in/ryan-stock-bb955555", "https://www.linkedin.com/in/ryan-stock-bb955555/"),
    ("communityplaytools.com", "https://www.communityplaytools.com/"),
]):
    if i:
        sep = p.add_run("  |  ")
        set_font(sep, "Arial", 8.4, MUTED)
    if url:
        add_hyperlink(p, label, url, color=PINE, underline=False)
    else:
        rr = p.add_run(label)
        set_font(rr, "Arial", 8.4, MUTED)

add_section_heading("Professional summary")
add_body(
    "PMP-certified senior program and innovation leader with 11+ years aligning cross-functional teams, redesigning delivery systems, and moving ideas from exploration through evidence to launch. Brings enterprise-scale PlayStation commerce leadership, confident facilitation, civic coalition experience, and founder-mode product delivery."
)

add_section_heading("Core strengths")
add_body(
    "Program & portfolio leadership  |  Release Train Engineering (SAFe)  |  Operating cadence & change  |  Cross-functional facilitation  |  Product discovery & prototyping  |  Payments & ecommerce  |  External partner delivery  |  Public speaking & workshops  |  Civic and stakeholder leadership",
    size=9.05,
    color=MUTED,
)

add_section_heading("Professional experience")
add_role(
    "Founder & Principal Consultant",
    "Community Play Tools",
    "Aug 2025 - Present",
    "St. Petersburg, FL",
    "Civic technology, participation systems, and live products",
)
add_bullet(
    "Produced St. Pete Detective Club, a paid six-location field mystery. Sold 20 seats, ran crews of 3 to 6 across a 2.5-mile route, and owned concept, software, physical production, venues, ticketing, promotion, and live hosting."
)
add_bullet(
    "Designed, built, and deployed the custom React/TypeScript and Supabase platform, including shared crew state, progressive evidence unlocks, and an interactive endgame, then ran it under live player load."
)
add_bullet(
    "Designed and shipped Signal Fire, a Rails and Expo place-based event product; designed a six-format Williams Park engagement suite and a reusable perception-baseline methodology."
)

add_role(
    "Senior Program Manager",
    "Sony Interactive Entertainment, PlayStation Network Commerce",
    "Dec 2017 - Jul 2025",
    "San Diego, CA (remote from 2020)",
    "Global ecommerce and payments platform serving 110M+ monthly active users",
)
add_subrole("Release Train Engineer, Commerce Agile Release Train")
add_bullet(
    "Served as facilitator and servant leader for 6+ Agile teams across regions, currencies, payment systems, and regulatory domains; led PI Planning, ART Sync, System Demos, Inspect & Adapt, risk, and dependency management."
)
add_bullet(
    "Replaced quarterly SAFe planning with a custom monthly cadence and led pandemic-era remote-first change; pioneered “remote tables,” later adopted as organizational standard practice."
)
add_bullet(
    "Led the Rally-to-Jira migration and restructured the product feature funnel with Product Management, improving prioritization and work-in-progress transparency."
)
add_subrole("Producer & Innovation Lead, PlayStation Direct")
add_bullet(
    "Produced Engineering Excellence hackathons for 80 to 240 participants across 6+ engineering teams, generating 30+ prototypes per event; one production adoption reduced cloud infrastructure costs 40%."
)
add_bullet(
    "Partnered with IDEO on AR, VR, and tactile low-screen concepts and used focus-group variant testing to determine which prototypes advanced."
)
add_subrole("Commerce Program Manager, Payments, Tax & Fraud")
add_bullet(
    "Led the PlayStation Direct PS5 preorder and prerelease checkout flow behind 4.5 million units sold in the launch quarter, coordinating engineering, product, finance, legal, and external partners."
)
add_bullet(
    "Delivered multi-region payment, wallet, and tax programs and served as primary delivery liaison with PayPal, Worldpay, Klarna, Chase, and Adyen, owning timelines, SOWs, and service obligations where applicable."
)

add_role(
    "Program Manager",
    "Omnitracs (formerly Qualcomm)",
    "Jun 2014 - Jan 2018",
    "Dallas, TX",
    "Enterprise fleet-management technology",
)
add_bullet(
    "Directed a next-generation fleet-management SaaS platform from concept to general availability, coordinating hardware, software, QA, procurement, sales, and legal to deliver the flagship app-enabling product on time and within budget."
)
add_bullet(
    "Led the Waterfall-to-Agile transition as Scrum Master, removing legacy process redundancies and improving release-cycle efficiency."
)

# Deliberate page balance: all enterprise experience stays on page one.
doc.add_page_break()

add_section_heading("Community & civic leadership")
add_role(
    "Board Member; Former President & Vice President",
    "Beautiful PB",
    "Feb 2023 - Present",
    "San Diego, CA",
    "President May 2024 - Jan 2026; Vice President 2023 - 2024",
)
add_bullet(
    "Led a 12-member volunteer board, created a three-pillar structure across Mobility, Arts, and Greening, built succession planning, and scaled the annual operating budget from under $1,000 to $40,000."
)
add_bullet(
    "Secured a $30,000 County of San Diego grant, then owned vendor selection and private-landowner siting to deploy three continuous automated counters, leading the 2023-2025 modernization of PB Counts from annual counting toward year-round monitoring."
)
add_bullet(
    "Built public data visualizations and added near-miss conflict data to the citizen-science program, translating raw sensor output into evidence for a human-safety policy conversation."
)
add_bullet(
    "Converted major-grant credibility into a $450,000 capital campaign for PB Arts Center restoration and served as primary spokesperson across city, county, nonprofit, and media stakeholders."
)
add_bullet(
    "Ran six workshops that moved contested PB Pathways Phase 3 to unanimous Planning Group recommendation and the top Parking Board approval result despite organized opposition."
)

add_role(
    "Board Member & Streets Chair",
    "Pacific Beach Community Planning Group",
    "2024 - 2025",
    "San Diego, CA",
    "City-recognized planning advisory group",
)
add_bullet(
    "Chaired the Streets & Sidewalks Subcommittee and presented a Garnet Avenue pedestrianization vision that ranked among the top three projects at the community capital-improvement fair."
)

add_section_heading("Selected product & innovation evidence")
add_bullet(
    "PlayStation innovation programs: structured experimentation, focus-group variant testing, 30+ prototypes per hackathon, and a clear path from exploration to production evidence."
)
add_bullet(
    "St. Pete Detective Club: paid, live, end-to-end product ownership across software, narrative, physical production, partners, promotion, and operations."
)
add_bullet(
    "Signal Fire: designed, built, and deployed place-based participation product; later deprioritized after evaluating the business model against available time and resources."
)
add_bullet(
    "Community Merit: designed, not shipped, a civic participation system based on check-ins, badges, and steward progression."
)

add_section_heading("Leadership approach")
add_bullet(
    "Align the real constraints: make decisions, risks, and dependencies visible enough for cross-functional teams to act without relying on private context."
)
add_bullet(
    "Prototype for evidence: use the lightest credible test, workshop, or live pilot that can replace abstract disagreement with observable behavior."
)
add_bullet(
    "Deliver through change: protect the outcome, retune the operating plan as conditions move, and remain accountable for what reaches the customer or community."
)

add_section_heading("Speaking & facilitation")
add_body(
    "Facilitated enterprise planning across 6+ Agile teams, produced innovation events for 80 to 240 participants, led six contested public workshops, represented a nonprofit before civic boards and media, and hosted a paid live field experience.",
    size=9.35,
)

add_section_heading("Education, certification & recognition")
add_body(
    "B.Sc., Biology & Bioengineering, San Diego State University, 2014  |  Project Management Professional (PMP), Project Management Institute, 2017 (PMI #2052114)",
    size=9.35,
)
add_body(
    "Select coursework: Creating & Sustaining Innovative Culture (University of Queensland), Game Theory (Yale), Cryptocurrency Engineering (MIT OpenCourseWare)",
    size=9.05,
    color=MUTED,
)
add_body(
    "Media: San Diego Union-Tribune, Times of San Diego, Beach & Bay Press  |  Eagle Scout  |  U.S. National Sprint Kayak Team (Under-23 and Pan American Team)",
    size=9.05,
    color=MUTED,
)

add_section_heading("Tools & delivery practice")
add_body(
    "Jira, Rally, Figma, React, TypeScript, React Native (Expo), Ruby on Rails, Supabase, Netlify. Accountable, AI-accelerated delivery: Ryan directs the work, reviews outputs, integrates systems, and retains accountability for what ships.",
    size=9.05,
    color=MUTED,
)

# Quiet page footer with real PAGE field.
for sec in doc.sections:
    footer = sec.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("RYAN STOCK  ·  ")
    set_font(r, "Arial", 7.5, MUTED, bold=True)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run = p.add_run()
    run._r.extend([fld_char1, instr, fld_char2])
    set_font(run, "Arial", 7.5, MUTED)

core = doc.core_properties
core.title = "Ryan Stock - Senior Program & Innovation Leader"
core.subject = "Public resume"
core.author = "Ryan Stock"
core.keywords = "program leadership, innovation, RTE, SAFe, civic leadership, product delivery"
core.comments = "Public role-neutral resume, updated July 2026."

doc.save(OUT)
print(OUT)
